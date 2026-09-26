from decimal import Decimal

from django.core.exceptions import ValidationError
from django.test import TestCase

from .models import Equipment, EquipmentHourRecord, Operator
from .services import validate_cumulative_hours


class CumulativeHourValidationTests(TestCase):
    def setUp(self):
        self.operator = Operator.objects.create(operator_id="TEST001")
        self.equipment = Equipment.objects.create(
            equipment_id="QA01", initial_hours=Decimal("110")
        )
        EquipmentHourRecord.objects.create(
            operator=self.operator,
            equipment=self.equipment,
            previous_hours=Decimal("110"),
            total_hours=Decimal("118"),
            shift_hours=Decimal("8"),
        )

    def test_accepts_eight_hours(self):
        bounds = validate_cumulative_hours(
            self.equipment, Decimal("126"), self.operator
        )
        self.assertEqual(bounds.previous, Decimal("118"))

    def test_accepts_same_reading(self):
        bounds = validate_cumulative_hours(
            self.equipment,
            Decimal("118"),
            Operator.objects.create(operator_id="TEST003"),
        )
        self.assertEqual(bounds.maximum, Decimal("126"))

    def test_accepts_partial_shift(self):
        validate_cumulative_hours(self.equipment, Decimal("124"), self.operator)

    def test_rejects_lower_reading(self):
        with self.assertRaisesMessage(
            ValidationError,
            "বর্তমান সময় এর চেয়ে কম হতে পারবে না",
        ):
            validate_cumulative_hours(self.equipment, Decimal("117"), self.operator)

    def test_rejects_more_than_eight_hours(self):
        with self.assertRaisesMessage(
            ValidationError,
            "সর্বোচ্চ গ্রহণযোগ্য সময় 126 ঘণ্টা",
        ):
            validate_cumulative_hours(self.equipment, Decimal("127"), self.operator)

    def test_accepts_decimal_readings(self):
        decimal_equipment = Equipment.objects.create(
            equipment_id="QA03", initial_hours=Decimal("118.5")
        )
        validate_cumulative_hours(
            decimal_equipment,
            Decimal("126.5"),
            operator=Operator.objects.create(operator_id="TEST002"),
        )

    def test_equipment_readings_are_independent(self):
        other = Equipment.objects.create(
            equipment_id="QA02", initial_hours=Decimal("90")
        )
        bounds = validate_cumulative_hours(other, Decimal("98"), self.operator)
        self.assertEqual(bounds.previous, Decimal("90"))


class EntryViewTests(TestCase):
    def test_entry_has_active_choices(self):
        Operator.objects.create(operator_id="TEST001", is_active=True)
        Operator.objects.create(operator_id="TEST002", is_active=False)
        Equipment.objects.create(equipment_id="QA01", initial_hours=Decimal("110"))
        response = self.client.get("/")
        self.assertContains(response, "TEST001")
        self.assertNotContains(response, "TEST002")


class AdminExcelExportTests(TestCase):
    def setUp(self):
        from io import BytesIO
        from django.contrib.auth.models import User
        from openpyxl import load_workbook

        self.BytesIO = BytesIO
        self.load_workbook = load_workbook
        self.admin_user = User.objects.create_superuser(
            username="admin_test", password="password123", email="admin@example.com"
        )
        self.operator = Operator.objects.create(
            operator_id="OP100", name="Rahim Ahmed"
        )
        self.equipment1 = Equipment.objects.create(
            equipment_id="EQ01", initial_hours=Decimal("100")
        )
        self.equipment2 = Equipment.objects.create(
            equipment_id="EQ02", initial_hours=Decimal("200")
        )
        self.record1 = EquipmentHourRecord.objects.create(
            operator=self.operator,
            equipment=self.equipment1,
            previous_hours=Decimal("100"),
            total_hours=Decimal("108"),
            shift_hours=Decimal("8"),
        )
        self.record2 = EquipmentHourRecord.objects.create(
            operator=self.operator,
            equipment=self.equipment2,
            previous_hours=Decimal("200"),
            total_hours=Decimal("205"),
            shift_hours=Decimal("5"),
        )

    def test_changelist_has_export_button(self):
        self.client.force_login(self.admin_user)
        response = self.client.get("/admin/equipment/equipmenthourrecord/")
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Export to Excel")
        self.assertContains(response, "export-excel/")

    def test_admin_export_excel_view(self):
        self.client.force_login(self.admin_user)
        response = self.client.get("/admin/equipment/equipmenthourrecord/export-excel/")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn("attachment; filename=", response["Content-Disposition"])

        wb = self.load_workbook(self.BytesIO(response.content))
        ws = wb.active
        self.assertEqual(ws.title, "Equipment Hours")
        self.assertEqual(ws.cell(row=1, column=1).value, "ID")
        self.assertEqual(ws.cell(row=1, column=4).value, "Operator Name")
        self.assertEqual(ws.cell(row=1, column=8).value, "Shift Hours")

        # Two records + 1 header + 1 total row = 4 rows
        self.assertEqual(ws.max_row, 4)
        self.assertEqual(ws.cell(row=4, column=1).value, "Total")
        self.assertEqual(ws.cell(row=4, column=8).value, "=SUM(H2:H3)")

    def test_admin_export_excel_with_filter(self):
        self.client.force_login(self.admin_user)
        response = self.client.get(
            f"/admin/equipment/equipmenthourrecord/export-excel/?equipment__id__exact={self.equipment1.pk}"
        )
        self.assertEqual(response.status_code, 200)
        wb = self.load_workbook(self.BytesIO(response.content))
        ws = wb.active
        # One record + 1 header + 1 total row = 3 rows
        self.assertEqual(ws.max_row, 3)
        self.assertEqual(ws.cell(row=2, column=5).value, "EQ01")

    def test_admin_export_action(self):
        self.client.force_login(self.admin_user)
        data = {
            "action": "export_to_excel",
            "_selected_action": [str(self.record1.pk)],
        }
        response = self.client.post("/admin/equipment/equipmenthourrecord/", data)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn("equipment_hours_selected_", response["Content-Disposition"])

        wb = self.load_workbook(self.BytesIO(response.content))
        ws = wb.active
        # Selected 1 record + 1 header + 1 total row = 3 rows
        self.assertEqual(ws.max_row, 3)
        self.assertEqual(ws.cell(row=2, column=1).value, self.record1.pk)

    def test_unauthenticated_export_denied(self):
        response = self.client.get("/admin/equipment/equipmenthourrecord/export-excel/")
        self.assertEqual(response.status_code, 302)

    def test_frontend_report_export(self):
        self.client.force_login(self.admin_user)
        response = self.client.get("/reports/export/")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )