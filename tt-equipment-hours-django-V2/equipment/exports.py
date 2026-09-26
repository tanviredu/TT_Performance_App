from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from django.http import HttpResponse
from django.utils import timezone


def export_equipment_hours_to_excel(
    queryset, filename_prefix: str = "equipment_hours"
) -> HttpResponse:
    """Generate an Excel (.xlsx) file for the given EquipmentHourRecord queryset."""
    records = queryset.select_related("operator", "equipment").order_by(
        "-created_at", "-id"
    )

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Equipment Hours"

    # Ensure grid lines are visible in Excel viewers
    worksheet.views.sheetView[0].showGridLines = True

    headers = [
        "ID",
        "Date",
        "Operator ID",
        "Operator Name",
        "Equipment ID",
        "Previous Hours",
        "Total Hours",
        "Shift Hours",
        "Created At",
    ]
    worksheet.append(headers)

    header_fill = PatternFill(
        start_color="1E3A8A", end_color="1E3A8A", fill_type="solid"
    )
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_alignment = Alignment(
        horizontal="center", vertical="center", wrap_text=True
    )

    thin_border_side = Side(style="thin", color="CBD5E1")
    cell_border = Border(
        left=thin_border_side,
        right=thin_border_side,
        top=thin_border_side,
        bottom=thin_border_side,
    )

    zebra_fill = PatternFill(
        start_color="F8FAFC", end_color="F8FAFC", fill_type="solid"
    )
    white_fill = PatternFill(
        start_color="FFFFFF", end_color="FFFFFF", fill_type="solid"
    )
    data_font = Font(name="Calibri", size=11)

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    # Format header row
    worksheet.row_dimensions[1].height = 26
    for col_idx in range(1, len(headers) + 1):
        cell = worksheet.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = cell_border

    tz = timezone.get_current_timezone()

    current_row = 2
    for record in records:
        created_at_dt = (
            record.created_at.astimezone(tz).replace(tzinfo=None)
            if record.created_at
            else None
        )
        operator_name = record.operator.name if record.operator else ""

        worksheet.append(
            [
                record.pk,
                record.date,
                record.operator.operator_id if record.operator else "",
                operator_name,
                record.equipment.equipment_id if record.equipment else "",
                float(record.previous_hours)
                if record.previous_hours is not None
                else 0.0,
                float(record.total_hours)
                if record.total_hours is not None
                else 0.0,
                float(record.shift_hours)
                if record.shift_hours is not None
                else 0.0,
                created_at_dt,
            ]
        )

        worksheet.row_dimensions[current_row].height = 20
        fill = zebra_fill if current_row % 2 == 0 else white_fill

        # Column formatting
        worksheet.cell(row=current_row, column=1).alignment = align_center
        worksheet.cell(row=current_row, column=1).number_format = "#,##0"

        worksheet.cell(row=current_row, column=2).alignment = align_center
        worksheet.cell(row=current_row, column=2).number_format = "YYYY-MM-DD"

        worksheet.cell(row=current_row, column=3).alignment = align_center
        worksheet.cell(row=current_row, column=4).alignment = align_left
        worksheet.cell(row=current_row, column=5).alignment = align_center

        worksheet.cell(row=current_row, column=6).alignment = align_right
        worksheet.cell(row=current_row, column=6).number_format = "#,##0.00"

        worksheet.cell(row=current_row, column=7).alignment = align_right
        worksheet.cell(row=current_row, column=7).number_format = "#,##0.00"

        worksheet.cell(row=current_row, column=8).alignment = align_right
        worksheet.cell(row=current_row, column=8).number_format = "#,##0.00"

        worksheet.cell(row=current_row, column=9).alignment = align_center
        worksheet.cell(row=current_row, column=9).number_format = (
            "YYYY-MM-DD HH:MM:SS"
        )

        for col_idx in range(1, len(headers) + 1):
            cell = worksheet.cell(row=current_row, column=col_idx)
            cell.font = data_font
            cell.fill = fill
            cell.border = cell_border

        current_row += 1

    total_records = current_row - 2

    # Add Summary / Total row if records exist
    if total_records > 0:
        summary_row = current_row
        worksheet.row_dimensions[summary_row].height = 22

        total_font = Font(name="Calibri", size=11, bold=True)
        double_bottom_border = Border(
            top=Side(style="thin", color="94A3B8"),
            bottom=Side(style="double", color="1E293B"),
            left=thin_border_side,
            right=thin_border_side,
        )
        total_fill = PatternFill(
            start_color="F1F5F9", end_color="F1F5F9", fill_type="solid"
        )

        label_cell = worksheet.cell(row=summary_row, column=1, value="Total")
        label_cell.alignment = align_center

        shift_sum_cell = worksheet.cell(
            row=summary_row,
            column=8,
            value=f"=SUM(H2:H{summary_row - 1})",
        )
        shift_sum_cell.alignment = align_right
        shift_sum_cell.number_format = "#,##0.00"

        for col_idx in range(1, len(headers) + 1):
            cell = worksheet.cell(row=summary_row, column=col_idx)
            cell.font = total_font
            cell.fill = total_fill
            cell.border = double_bottom_border

    # Freeze header row
    worksheet.freeze_panes = "A2"

    # Auto-filter
    last_row_filter = max(current_row - 1, 1)
    last_col_letter = get_column_letter(len(headers))
    worksheet.auto_filter.ref = f"A1:{last_col_letter}{last_row_filter}"

    # Auto-adjust column widths
    for col in worksheet.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val = cell.value
            if val is not None:
                if isinstance(val, float):
                    s = f"{val:,.2f}"
                else:
                    s = str(val)
                max_len = max(max_len, len(s))
        worksheet.column_dimensions[col_letter].width = max(max_len + 4, 12)

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    timestamp = timezone.localtime(timezone.now()).strftime("%Y%m%d_%H%M%S")
    filename = f"{filename_prefix}_{timestamp}.xlsx"

    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response
